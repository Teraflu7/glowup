from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "REPORT.docx"
ASSETS = ROOT / "report-assets"

INK = RGBColor(33, 32, 38)
MUTED = RGBColor(91, 86, 99)
PINK = "EAA0B2"
PALE_PINK = "FBECEF"
PALE_GRAY = "F5F3F6"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, bold=None, color=INK, italic=None):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = color


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size in (("Title", 30), ("Heading 1", 19), ("Heading 2", 14), ("Heading 3", 11.5)):
        style = document.styles[name]
        style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    caption = document.styles["Caption"]
    caption.font.name = "Aptos"
    caption._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Aptos")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("GlowUp  •  Front End Assignment 4  •  SE-2516  •  ")
    set_run_font(run, 8.5, color=MUTED)
    add_page_number(footer)


def add_title_page(document):
    top = document.add_paragraph()
    top.paragraph_format.space_before = Pt(30)
    top.paragraph_format.space_after = Pt(8)
    top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(top.add_run("GLOWUP"), 13, bold=True, color=RGBColor(151, 83, 105))

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(36)
    title.paragraph_format.space_after = Pt(8)
    set_run_font(title.add_run("Front End Assignment 4"), 30, bold=True, color=RGBColor(0, 0, 0))

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    set_run_font(subtitle.add_run("JavaScript interactions in the GlowUp website"), 15, color=MUTED)

    table = document.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    metadata = [
        ("Team", "Tamiris Mukieva and Alina Abdubek"),
        ("Group", "SE-2516"),
        ("Repository", "github.com/Teraflu7/glowup"),
        ("Deployment", "teraflu7.github.io/glowup/"),
        ("Technology", "HTML5, CSS3, Bootstrap 5.3.8, vanilla JavaScript"),
    ]
    for index, (label, value) in enumerate(metadata):
        left, right = table.rows[index].cells
        left.width = Inches(1.75)
        right.width = Inches(4.9)
        set_cell_shading(left, PALE_PINK)
        set_cell_shading(right, "FFFFFF" if index % 2 == 0 else PALE_GRAY)
        for cell in (left, right):
            set_cell_margins(cell, top=105, bottom=105)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        left.text = label
        right.text = value
        set_run_font(left.paragraphs[0].runs[0], 10, bold=True)
        set_run_font(right.paragraphs[0].runs[0], 10)

    summary = document.add_paragraph()
    summary.paragraph_format.space_before = Pt(26)
    summary.paragraph_format.space_after = Pt(0)
    summary.paragraph_format.left_indent = Inches(0.38)
    summary.paragraph_format.right_indent = Inches(0.38)
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        summary.add_run(
            "This report documents five required JavaScript features: custom form validation, "
            "an FAQ accordion, a popup form, background-color controls, and a live date and time display."
        ),
        11.5,
        color=MUTED,
    )

    year = document.add_paragraph()
    year.paragraph_format.space_before = Pt(42)
    year.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(year.add_run("October 2026"), 10, bold=True, color=MUTED)
    document.add_page_break()


def add_intro(document):
    document.add_heading("1. Project objective and result", level=1)
    document.add_paragraph(
        "The objective was to transform the existing responsive GlowUp website into an interactive "
        "multi-page experience using JavaScript. The completed version preserves the original visual "
        "design while adding accessible controls, clear form feedback, and responsive behavior."
    )

    document.add_heading("Requirement coverage", level=2)
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(("Requirement", "Implementation", "Main files", "Owner")):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, PINK)
        set_cell_margins(cell)
        set_run_font(cell.paragraphs[0].runs[0], 9, bold=True)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("Form validation", "Required values, email format, password length and match", "about.html; js/signup.js", "Alina"),
        ("FAQ accordion", "Hidden answers, smooth open/close, one item at a time", "faq.html; js/faq.js", "Tamiris"),
        ("Popup form", "Native dialog, close button, backdrop, Escape", "about.html; css/signup-dialog.css", "Alina"),
        ("Background color", "Four-color cycle and reset with status text", "index.html; js/home.js", "Tamiris"),
        ("Date and time", "Localized display updated every second", "index.html; js/home.js", "Tamiris"),
    ]
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = value
            set_cell_margins(cells[index])
            if row_index % 2:
                set_cell_shading(cells[index], PALE_GRAY)
            set_run_font(cells[index].paragraphs[0].runs[0], 8.5)

    document.add_heading("Integration approach", level=2)
    for item in (
        "Each page loads its JavaScript module with defer, after the shared and component-specific CSS.",
        "Native HTML controls provide a reliable keyboard baseline; JavaScript adds state and feedback.",
        "ARIA attributes expose accordion state, validation state, and live status information.",
        "The features were tested locally and then verified from the deployed GitHub Pages site.",
    ):
        document.add_paragraph(item, style="List Bullet")


def add_figure(document, filename, caption, width=6.65):
    path = ASSETS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.add_run(caption)


def add_validation_and_popup(document):
    document.add_page_break()
    document.add_heading("2. Form validation and popup", level=1)
    document.add_paragraph(
        "Alina implemented the contact interaction in about.html. The form is placed inside a native "
        "dialog that is closed by default. The opening button calls showModal(), while the close button, "
        "backdrop selection, and Escape key all dismiss it. Focus returns to the opening control."
    )
    document.add_heading("Validation logic", level=2)
    for item in (
        "All fields are required and text values are trimmed before checking.",
        "Email is checked against a readable email pattern.",
        "The password must contain at least eight characters, and confirmation must match exactly.",
        "Invalid controls receive aria-invalid, an inline error message, and focus on the first problem.",
        "Success feedback appears only when every rule passes.",
    ):
        document.add_paragraph(item, style="List Bullet")
    add_figure(
        document,
        "assignment4-validation-errors.png",
        "Figure 1. Inline validation messages identify every missing or invalid field.",
    )
    add_figure(
        document,
        "assignment4-popup-open.png",
        "Figure 2. The responsive contact dialog opens above the About page.",
    )


def add_accordion(document):
    document.add_page_break()
    document.add_heading("3. FAQ accordion", level=1)
    document.add_paragraph(
        "Tamiris converted the four FAQ questions into native buttons. All answers are hidden when the "
        "script initializes. Selecting a question expands its connected panel with a smooth transition; "
        "selecting it again closes it, and opening a new question closes the previous answer."
    )
    document.add_heading("Accessibility details", level=2)
    for item in (
        "aria-controls connects each trigger to its answer panel.",
        "aria-expanded reports the current state to assistive technology.",
        "Native buttons support Tab, Enter, and Space without custom key emulation.",
        "Reduced-motion preferences disable nonessential animation.",
    ):
        document.add_paragraph(item, style="List Bullet")
    add_figure(
        document,
        "assignment4-accordion-open.png",
        "Figure 3. One FAQ answer is expanded while the other answers remain collapsed.",
    )


def add_home_interactions(document):
    document.add_page_break()
    document.add_heading("4. Background controls and live clock", level=1)
    intro = document.add_paragraph(
        "Tamiris added an interaction panel to the home page. The background control cycles through four "
        "soft colors by updating the existing --bg CSS custom property. A separate reset action restores "
        "the first theme, and a live text message reports the selected color."
    )
    intro.paragraph_format.space_before = Pt(8)
    add_figure(
        document,
        "assignment4-background-color.png",
        "Figure 4. The home page after changing to another background theme.",
    )
    document.add_paragraph(
        "The same module creates a Date object, formats the value with Intl.DateTimeFormat, writes it to a "
        "semantic time element, and updates the display every second. The machine-readable datetime "
        "attribute remains synchronized with the visible value."
    )
    add_figure(
        document,
        "assignment4-date-time.png",
        "Figure 5. The localized current date and time display updates without reloading the page.",
    )


def add_responsive_and_testing(document):
    document.add_page_break()
    document.add_heading("5. Responsive behavior and verification", level=1)
    document.add_paragraph(
        "The interactive components adapt to narrow screens without horizontal scrolling. The dialog uses "
        "a bounded viewport width and internal scrolling when necessary, while buttons and messages remain "
        "large enough for touch interaction."
    )
    add_figure(
        document,
        "assignment4-mobile-view.png",
        "Figure 6. The popup form remains usable at a 375 px mobile viewport.",
        width=3.15,
    )

    document.add_page_break()
    document.add_heading("Functional test results", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(("Check", "Evidence", "Result")):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, PINK)
        set_cell_margins(cell)
        set_run_font(cell.paragraphs[0].runs[0], 9, bold=True)
    set_repeat_table_header(table.rows[0])
    tests = [
        ("Form rules", "Empty, malformed email, short password, mismatch, valid values", "Passed"),
        ("Dialog closing", "Close button, backdrop, Escape; focus restoration", "Passed"),
        ("Accordion", "Hidden initially; open, close, single-open behavior", "Passed"),
        ("Home controls", "Four colors, reset, visible status, one-second clock update", "Passed"),
        ("Responsive", "375 px, 768 px, and 1440 px", "Passed"),
        ("Browsers", "Google Chrome and Microsoft Edge", "Passed"),
        ("Console/network", "No JavaScript errors; local resources return successfully", "Passed"),
        ("Deployment", "GitHub Pages workflow and public feature resources", "Passed"),
    ]
    for row_index, test in enumerate(tests):
        cells = table.add_row().cells
        for index, value in enumerate(test):
            cells[index].text = value
            set_cell_margins(cells[index])
            if row_index % 2:
                set_cell_shading(cells[index], PALE_GRAY)
            set_run_font(cells[index].paragraphs[0].runs[0], 8.5, bold=(index == 2))


def add_team_and_conclusion(document):
    document.add_page_break()
    document.add_heading("6. Team contribution", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(("Member", "Primary implementation", "Files")):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, PINK)
        set_cell_margins(cell)
        set_run_font(cell.paragraphs[0].runs[0], 9, bold=True)
    contributions = [
        (
            "Tamiris Mukieva",
            "FAQ accordion; background color changer; live date and time",
            "index.html, faq.html, js/home.js, js/faq.js, related CSS",
        ),
        (
            "Alina Abdubek",
            "Popup contact form; custom form validation; inline feedback",
            "about.html, js/signup.js, css/signup-dialog.css",
        ),
        (
            "Both members",
            "Integration review, cross-browser checks, screenshots, deployment and defense preparation",
            "README.md, REPORT.md, REPORT.docx, report-assets/",
        ),
    ]
    for row_index, contribution in enumerate(contributions):
        cells = table.add_row().cells
        for index, value in enumerate(contribution):
            cells[index].text = value
            set_cell_margins(cells[index])
            if row_index % 2:
                set_cell_shading(cells[index], PALE_GRAY)
            set_run_font(cells[index].paragraphs[0].runs[0], 9)

    document.add_heading("7. Deployment", level=1)
    document.add_paragraph(
        "The shared main branch is published through GitHub Pages at "
        "https://teraflu7.github.io/glowup/. The successful deployment contains the same home, FAQ, and "
        "signup JavaScript modules documented in this report."
    )

    document.add_heading("8. Reflection and conclusion", level=1)
    document.add_paragraph(
        "Assignment 4 changed GlowUp from a mostly static responsive website into an interactive application. "
        "The team practiced DOM selection, event listeners, state changes, validation, the Date API, focus "
        "management, accessible feedback, and organization of JavaScript by page."
    )
    document.add_paragraph(
        "All five required features are implemented, responsive, documented, and deployed. The repository "
        "history records contributions from both team accounts, and the project is ready for submission and defense."
    )

    box = document.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.cell(0, 0)
    set_cell_shading(cell, PALE_PINK)
    set_cell_margins(cell, top=170, start=190, bottom=170, end=190)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        paragraph.add_run("Final result: all Assignment 4 requirements completed and verified"),
        11,
        bold=True,
    )


def build_report():
    document = Document()
    configure_document(document)
    add_title_page(document)
    add_intro(document)
    add_validation_and_popup(document)
    add_accordion(document)
    add_home_interactions(document)
    add_responsive_and_testing(document)
    add_team_and_conclusion(document)

    properties = document.core_properties
    properties.title = "Front End Assignment 4"
    properties.subject = "JavaScript interactions in the GlowUp website"
    properties.author = "Tamiris Mukieva; Alina Abdubek"
    properties.keywords = "JavaScript, form validation, accordion, dialog, date and time, GlowUp"

    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_report()
